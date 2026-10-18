"""Etapa 5 — Genera el informe .docx del trabajo práctico.

Uso:
    python src/generar_informe.py --modulo N --fecha DD/MM/AAAA [--salida RUTA]

Todos los números salen de outputs/resultados.json, outputs/historial.json y los demás archivos de
outputs/; las figuras y tablas se numeran por orden de aparición. El formato replica el del TP de
referencia (TP Árboles de Decisión), medido sobre su PDF: márgenes de 2,5 cm, encabezado con el logo UP,
"Facultad de Ingeniería" y una barra de 3 pt, Arial 11 justificado con interlineado 1,5, tablas con
encabezado gris y recuadros grises para código y resultados.
"""

import argparse
import ast
import math
import json
import re
import sys
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAIZ = Path(__file__).resolve().parent.parent
SALIDAS = RAIZ / "outputs"
FIGURAS = SALIDAS / "figuras"
SRC = RAIZ / "src"
LOGO = RAIZ / "recursos" / "logo_up.png"
SALIDA_POR_DEFECTO = RAIZ / "entrega" / "TP_Redes_Neuronales_Convolucionales_-_Zaragoza_Yago.docx"

# Formato medido sobre el PDF de referencia
FUENTE, MONO, SIMBOLOS = "Arial", "Courier New", "Segoe UI Symbol"
GRIS_ENCABEZADO, GRIS_CAJA, GRIS_CODIGO, BORDE_CODIGO = "D9D9D9", "F2F2F2", "F5F5F5", "BFBFBF"
GRIS_EPIGRAFE = RGBColor(0x40, 0x40, 0x40)
VERDE, ROJO = RGBColor(0x1A, 0x7F, 0x37), RGBColor(0xC6, 0x28, 0x28)
ANCHO_TEXTO_CM = 16.0
ANCHO_TEXTO_PT = ANCHO_TEXTO_CM / 2.54 * 72
TABLAS_COMPACTAS = {"ejemplos", "epocas"}  # tablas largas de los anexos, con menos relleno vertical


# ---------------------------------------------------------------- formato de números

def n(valor, decimales=0):
    """Número con punto de miles y coma decimal."""
    s = f"{valor:,.{decimales}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(fraccion, decimales=2):
    """Fracción expresada como porcentaje."""
    return f"{n(100 * fraccion, decimales)} %"


def pct_fuente(fraccion):
    """Porcentaje con la precisión de la fuente (0.897 → 89,7 %; 0.8833 → 88,33 %), máximo 2 decimales."""
    decimales = len(repr(float(fraccion)).split(".")[1]) - 2
    return pct(fraccion, max(0, min(2, decimales)))


def dim(*valores):
    return " × ".join(n(v) for v in valores)


PALABRAS = {1: "una", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco", 6: "seis", 7: "siete", 8: "ocho",
            9: "nueve", 10: "diez", 11: "once", 12: "doce"}


def palabra(entero, femenino=True):
    texto = PALABRAS.get(entero, n(entero))
    return texto if femenino or entero != 1 else "uno"


PRENDAS_FEMENINAS = {"remera/top", "camisa", "sandalia", "zapatilla", "bota"}


def cantidad_prenda(cantidad, prenda):
    """'un abrigo', 'una camisa', 'dos remeras/tops', 'dos pantalones' (prenda en minúscula)."""
    if cantidad == 1:
        return f"{'una' if prenda in PRENDAS_FEMENINAS else 'un'} {prenda}"
    plural = "/".join(p[:-2] + "ones" if p.endswith("ón") else p + "s" for p in prenda.split("/"))
    return f"{palabra(cantidad)} {plural}"


# ---------------------------------------------------------------- utilidades de XML

def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _bordes_celda(celda, **lados):
    """lados: top/left/bottom/right = (tamaño en octavos de punto, color) o None para sin borde."""
    tcPr = celda._tc.get_or_add_tcPr()
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        valor = lados.get(lado)
        if valor is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(valor[0]))
            el.set(qn("w:color"), valor[1])
            el.set(qn("w:space"), "0")
        bordes.append(el)
    tcPr.append(bordes)


def _bordes_tabla(tabla, tam=8, color="000000"):
    tblPr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        if tam:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(tam))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        bordes.append(el)
    tblPr.append(bordes)


def _margenes_celdas(tabla, arriba=40, izquierda=100, abajo=40, derecha=100):
    """Márgenes internos de las celdas, en veinteavos de punto."""
    tblPr = tabla._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", arriba), ("left", izquierda), ("bottom", abajo), ("right", derecha)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(valor))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _anchos(tabla, anchos_cm):
    """Fija el ancho de cada columna en la grilla y en cada celda (LibreOffice usa la grilla)."""
    for j, ancho in enumerate(anchos_cm):
        tabla.columns[j].width = Cm(ancho)
        for celda in tabla.columns[j].cells:
            celda.width = Cm(ancho)


def _fila_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_partir(fila):
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _fuente(run, nombre=FUENTE, tam=11, negrita=None, cursiva=None, subrayado=None, color=None):
    run.font.name = nombre
    run.font.size = Pt(tam)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(atributo), nombre)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva
    if subrayado is not None:
        run.underline = subrayado
    if color is not None:
        run.font.color.rgb = color


def _parrafo(par, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, interlineado=1.5, antes=0, despues=9,
             juntar=False, sangria=None):
    pf = par.paragraph_format
    pf.alignment = alineacion
    pf.line_spacing = interlineado
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.keep_with_next = juntar
    pf.widow_control = True
    if sangria is not None:
        pf.left_indent = Cm(sangria)


def _texto_con_formato(par, texto, tam=11, color=None):
    """Agrega texto interpretando **negrita** y _cursiva_."""
    for trozo in re.split(r"(\*\*.+?\*\*|(?<!\w)_[^_]+?_(?!\w))", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            _fuente(par.add_run(trozo[2:-2]), tam=tam, negrita=True, color=color)
        elif trozo.startswith("_") and trozo.endswith("_"):
            _fuente(par.add_run(trozo[1:-1]), tam=tam, cursiva=True, color=color)
        else:
            _fuente(par.add_run(trozo), tam=tam, color=color)


# ---------------------------------------------------------------- piezas del documento

def configurar_documento(doc):
    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.font.size = Pt(11)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.top_margin = Cm(5.0)          # el cuerpo empieza debajo del encabezado, como en la referencia
    sec.header_distance = Cm(1.27)    # el logo arranca a 36 pt del borde superior
    sec.footer_distance = Cm(1.27)
    encabezado(sec.header)


def encabezado(header):
    """Logo UP de 2,54 cm, 'Facultad de Ingeniería' en Arial 11 negrita y barra de 3 pt."""
    p = header.paragraphs[0]
    _parrafo(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 2)
    p.add_run().add_picture(str(LOGO), width=Cm(2.54))
    tabla = header.add_table(1, 2, width=Cm(ANCHO_TEXTO_CM))
    tabla.autofit = False
    _bordes_tabla(tabla, 0)
    _margenes_celdas(tabla, 0, 0, 0, 0)
    izq, der = tabla.rows[0].cells
    _anchos(tabla, [4.41, ANCHO_TEXTO_CM - 4.41])
    _bordes_celda(izq)
    _bordes_celda(der, bottom=(24, "000000"))
    for celda in (izq, der):
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        _parrafo(celda.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
    _fuente(izq.paragraphs[0].add_run("Facultad de Ingeniería"), negrita=True)
    final = header.add_paragraph()
    _parrafo(final, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
    _fuente(final.add_run(""), tam=2)


def portada(doc, modulo, actividad, fecha):
    """Recuadro de portada con la misma estructura y medidas que la referencia."""
    separador = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    _parrafo(separador, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 14, 0)
    caja = doc.add_table(1, 1)
    caja.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes_tabla(caja, 8)
    fila = caja.rows[0]
    fila.height, fila.height_rule = Pt(587), WD_ROW_HEIGHT_RULE.EXACTLY
    celda = fila.cells[0]
    celda.width = Cm(ANCHO_TEXTO_CM)

    def titulo(par, texto, tam, antes, despues, interlineado=1.0):
        _parrafo(par, WD_ALIGN_PARAGRAPH.CENTER, interlineado, antes, despues)
        _fuente(par.add_run(texto), tam=tam, negrita=True, cursiva=True, subrayado=True)

    titulo(celda.paragraphs[0], "TRABAJO PRÁCTICO", 20, 34, 0)
    titulo(celda.add_paragraph(), "Inteligencia Artificial", 18, 36, 0, Pt(29))
    titulo(celda.add_paragraph(), f"Módulo {modulo}", 18, 0, 0, Pt(29))
    titulo(celda.add_paragraph(), f"Actividad: {actividad}", 18, 0, 42, Pt(29))

    datos = celda.add_table(2, 3)
    datos.alignment = WD_TABLE_ALIGNMENT.CENTER
    datos.autofit = False
    _bordes_tabla(datos, 0)
    anchos = (Cm(2.65), Cm(5.48), Cm(6.70))
    filas = (("Legajo", "Apellido y Nombre", "e-mail"), ("0129980", "Zaragoza Yago", "yagozaragoza04@gmail.com"))
    for i, valores in enumerate(filas):
        datos.rows[i].height, datos.rows[i].height_rule = Pt(19.5), WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, valor in enumerate(valores):
            c = datos.rows[i].cells[j]
            c.width = anchos[j]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                _sombrear(c, GRIS_ENCABEZADO)
            p = c.paragraphs[0]
            _parrafo(p, WD_ALIGN_PARAGRAPH.CENTER if j == 2 else WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
            _fuente(p.add_run(valor), negrita=True)
    _anchos(datos, [2.65, 5.48, 6.70])

    for k, (etiqueta, valor) in enumerate((("Carrera", "Ingeniería Informática"), ("Modalidad", "Asincrónico"),
                                           ("Docente", "Ignacio Casais"), ("Fecha de entrega", fecha))):
        p = celda.add_paragraph()
        _parrafo(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 118 if k == 0 else 15, 0, sangria=0.25)
        _fuente(p.add_run(etiqueta), negrita=True, subrayado=True)
        _fuente(p.add_run(f": {valor}"), negrita=True)


def titulo1(doc, texto, salto=False):
    p = doc.add_paragraph()
    _parrafo(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 18, 12, juntar=True)
    p.paragraph_format.page_break_before = salto
    _fuente(p.add_run(texto), tam=14, negrita=True)


def titulo2(doc, texto):
    p = doc.add_paragraph()
    _parrafo(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 12, 6, juntar=True)
    _fuente(p.add_run(texto), tam=12, negrita=True)


def cuerpo(doc, texto, sangria=None, juntar=False):
    p = doc.add_paragraph()
    _parrafo(p, sangria=sangria, juntar=juntar)
    _texto_con_formato(p, texto)


def figura(doc, archivo, numero, epigrafe, ancho_max, alto_max):
    with Image.open(archivo) as im:
        w, h = im.size
    ancho = min(ancho_max, alto_max * w / h, ANCHO_TEXTO_CM)
    p = doc.add_paragraph()
    _parrafo(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6, 3, juntar=True)
    p.add_run().add_picture(str(archivo), width=Cm(ancho))
    e = doc.add_paragraph()
    _parrafo(e, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 12)
    _fuente(e.add_run(f"{numero}. {epigrafe}"), tam=9, cursiva=True, color=GRIS_EPIGRAFE)


def tabla(doc, numero, titulo, encabezados, filas, anchos, alineaciones=None, tam=10, resaltar=None, compacta=False):
    t = doc.add_paragraph()
    _parrafo(t, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6, 4, juntar=True)
    _fuente(t.add_run(f"{numero}. {titulo}"), tam=9, cursiva=True, color=GRIS_EPIGRAFE)
    alineaciones = alineaciones or ["l"] + ["c"] * (len(encabezados) - 1)
    resaltar = resaltar or {}
    mapa = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
    tab = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    _bordes_tabla(tab, 8)
    _margenes_celdas(tab, 8 if compacta else 30, 90, 8 if compacta else 30, 90)
    for i, fila in enumerate([encabezados] + filas):
        row = tab.rows[i]
        _no_partir(row)
        if i == 0:
            _fila_encabezado(row)
        estilo = resaltar.get(i - 1, {}) if i else {}
        for j, valor in enumerate(fila):
            c = row.cells[j]
            c.width = Cm(anchos[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                _sombrear(c, GRIS_ENCABEZADO)
            elif estilo.get("fondo"):
                _sombrear(c, estilo["fondo"])
            p = c.paragraphs[0]
            _parrafo(p, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else mapa[alineaciones[j]], 1.0, 0, 0)
            if isinstance(valor, tuple):  # (texto, color, fuente) para símbolos
                _fuente(p.add_run(valor[0]), nombre=valor[2] if len(valor) > 2 else FUENTE, tam=tam,
                        negrita=True, color=valor[1])
            else:
                _fuente(p.add_run(str(valor)), tam=tam, negrita=(i == 0) or estilo.get("negrita", False),
                        cursiva=estilo.get("cursiva", False))
    _anchos(tab, anchos)
    despues = doc.add_paragraph()
    _parrafo(despues, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 6)


def codigo(doc, texto):
    """Recuadro gris con borde fino y fuente monoespaciada; el tamaño se ajusta para no partir líneas."""
    lineas = texto.rstrip("\n").split("\n")
    largo = max(len(l) for l in lineas)
    tam = min(9.5, (ANCHO_TEXTO_PT - 14) / (0.6 * largo))
    tam = max(7.0, round(tam * 2) / 2 - 0.5 if tam < 9.5 else 9.5)
    t = doc.add_table(1, 1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _margenes_celdas(t, 100, 110, 100, 110)
    _no_partir(t.rows[0])  # el recuadro de código no se corta entre páginas
    c = t.rows[0].cells[0]
    c.width = Cm(ANCHO_TEXTO_CM)
    _sombrear(c, GRIS_CODIGO)
    _bordes_celda(c, top=(8, BORDE_CODIGO), left=(8, BORDE_CODIGO), bottom=(8, BORDE_CODIGO), right=(8, BORDE_CODIGO))
    for k, linea in enumerate(lineas):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        _parrafo(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
        _fuente(p.add_run(linea), nombre=MONO, tam=tam)
    despues = doc.add_paragraph()
    _parrafo(despues, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 6)


def caja(doc, titulo, parrafos):
    """Recuadro de resultado o advertencia: fondo gris, borde izquierdo de 2 pt."""
    t = doc.add_table(1, 1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _margenes_celdas(t, 100, 140, 100, 140)
    _no_partir(t.rows[0])
    c = t.rows[0].cells[0]
    c.width = Cm(ANCHO_TEXTO_CM)
    _sombrear(c, GRIS_CAJA)
    _bordes_celda(c, top=(8, "000000"), left=(16, "000000"), bottom=(8, "000000"), right=(8, "000000"))
    p = c.paragraphs[0]
    _parrafo(p, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0, 4, juntar=True)
    _fuente(p.add_run(titulo), negrita=True)
    for texto in parrafos:
        q = c.add_paragraph()
        _parrafo(q, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 0, 4)
        _texto_con_formato(q, texto, tam=10.5)
    despues = doc.add_paragraph()
    _parrafo(despues, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 6)


# ---------------------------------------------------------------- contenido

def fuente_python(ruta, funcion):
    """Código fuente de una función, tal como está en el archivo."""
    texto = ruta.read_text(encoding="utf-8")
    for nodo in ast.walk(ast.parse(texto)):
        if isinstance(nodo, ast.FunctionDef) and nodo.name == funcion:
            return ast.get_source_segment(texto, nodo)
    raise ValueError(f"no se encontró {funcion} en {ruta}")


def fragmento_fit(ruta):
    """Callbacks y llamada a fit de src/entrenar.py."""
    lineas = ruta.read_text(encoding="utf-8").split("\n")
    ini = next(i for i, l in enumerate(lineas) if l.strip() == "if args.prueba:")
    fin = next(i for i, l in enumerate(lineas) if i > ini and "callbacks=callbacks, verbose=2)" in l)
    return textwrap.dedent("\n".join(lineas[ini:fin + 1]))


def construir_bloques(R, H, B, E, entorno, args):
    """Arma la lista de bloques del informe. Las referencias [[fig:x]] y [[tab:x]] se resuelven después."""
    D, A, T, V, M = R["dataset"], R["arquitectura"], R["entrenamiento"], R["validacion"], R["mlp_referencia"]
    capas = {c["nombre"]: c for c in A["capas"]}
    cfg = {k: c.get("config", {}) for k, c in capas.items()}
    hp, part = T["hiperparametros"], T["particion"]
    mejor, previa = T["mejor_epoca"], T["corrida_previa"]
    g, ej, err = V["global"], V["ejemplos"], V["errores_alta_confianza"]
    bench, comp = V["benchmark"], V["comparacion_cnn_mlp"]
    kernel = dim(*cfg["conv1_1"]["kernel"])
    pool = cfg["pool1"]["pool"]
    convs = [c for c in A["capas"] if c["tipo"] == "Conv2D"]
    forma = lambda nombre: dim(*capas[nombre]["forma_salida"])
    n_por_clase_tr = set(D["distribucion_entrenamiento"].values())
    n_por_clase_te = set(D["distribucion_prueba"].values())
    assert len(n_por_clase_tr) == 1 and len(n_por_clase_te) == 1
    diag = D["diagnostico_csv_vs_idx"]

    # Afirmaciones del texto que dependen de los datos: si dejan de ser ciertas, el informe no se genera
    assert T["corte_por_early_stopping"] and mejor < T["epocas_ejecutadas"]
    assert all(e > mejor for e in T["epocas_con_accuracy_train_mayor_que_val"])
    assert all(V["confusion_prendas_de_torso"]["top3_dentro_del_grupo"])
    assert previa["mejor_epoca"] == previa["epocas_ejecutadas"] == previa["epocas_maximas"]
    assert g["accuracy_test"] < g["accuracy_validacion_mejor_epoca"]

    b = []
    add = b.append

    # Consigna (transcripción textual)
    add(("h1", "Consigna", {"salto": True}))
    add(("p", "Construir un modelo e implementar utilizando CNN. Para ello debés:"))
    for item in ("a. Buscar y seleccionar un conjunto de datos para trabajar (Kaggle / Google Dataset Search).",
                 "b. Programar (construir) la red neuronal.",
                 "c. Realizar el proceso de aprendizaje.",
                 "d. Realizar la validación del modelo utilizando ejemplos y cotejando el resultado."):
        add(("item", item))

    # 1. Introducción
    add(("h1", "1. Introducción y objetivo", {}))
    add(("intro", "Una red neuronal convolucional (CNN) es una red neuronal que usa la convolución, un filtro o "
         "kernel que recorre la imagen, en lugar de la multiplicación matricial general en al menos una de sus "
         "capas, y la combina con capas ReLU, de agrupación (pooling), flatten y totalmente conectadas "
         "(material de cátedra: PPT Redes Neuronales Convolucionales). El objetivo de este trabajo es "
         f"construir, entrenar y validar una CNN que clasifique imágenes de prendas del conjunto Fashion-MNIST "
         f"en sus {D['num_clases']} clases."))
    add(("intro", "El informe sigue el orden de la consigna. La sección 2 (consigna a) justifica la selección "
         "del conjunto de datos y describe su preprocesamiento; la sección 3 (consigna b) presenta la "
         "arquitectura de la red; la sección 4 (consigna c) explica el proceso de aprendizaje, y la sección 5 (consigna d) valida el modelo sobre el "
         "conjunto de prueba oficial con métricas por clase, ejemplos cotejados, el análisis de los errores y "
         "la comparación con una red sin convolución y con resultados publicados. Los anexos documentan el "
         "entorno y la ejecución ([[anx:entorno]]), los ejemplos cotejados ([[anx:ejemplos]]), el detalle por "
         "época ([[anx:epocas]]) y la verificación de los datos ([[anx:verificacion]])."))

    # 2. Consigna a
    add(("h1", "2. Consigna a: selección del conjunto de datos", {}))
    add(("h2", "Búsqueda y criterios"))
    add(("p", "La búsqueda y la descarga se realizaron en Kaggle, con seis criterios: que los datos fueran "
         "imágenes, insumo natural de una red neuronal convolucional; que el etiquetado fuera confiable y hecho "
         "por expertos; que las clases estuvieran balanceadas; que el tamaño permitiera entrenar sin GPU; que "
         "la licencia fuera abierta, y que existiera un benchmark publicado contra el cual cotejar los "
         "resultados."))
    add(("p", f"Se eligió Fashion-MNIST, de Zalando Research, publicado con licencia {D['licencia']} en Kaggle y "
         "en GitHub (repositorio oficial de Fashion-MNIST). Sus etiquetas provienen "
         "del código de silueta de cada producto, que fue etiquetado manualmente por expertos en moda de Zalando "
         "y revisado por otro equipo de la empresa (Xiao, Rasul y Vollgraf, 2017). Se descartó el MNIST original "
         "de dígitos manuscritos porque es el ejemplo que desarrolla el video de la cátedra (material de "
         "cátedra: video M9); Fashion-MNIST fue creado como reemplazo directo y más exigente de MNIST y conserva "
         "su formato, por lo que la arquitectura vista en clase se aplica sin cambios a un problema menos "
         "trivial."))
    add(("h2", "Descripción del conjunto"))
    lado = D["forma_imagen"][0]
    add(("p", f"El conjunto contiene {n(D['n_total'])} imágenes de prendas y accesorios en escala de grises de "
         f"{dim(lado, lado)} píxeles: {n(D['n_entrenamiento'])} de entrenamiento y {n(D['n_prueba'])} de prueba, "
         f"repartidas en {D['num_clases']} clases ([[tab:clases]]). Las clases están balanceadas: "
         f"{n(n_por_clase_tr.pop())} imágenes por clase en entrenamiento y {n(n_por_clase_te.pop())} en prueba "
         "([[fig:distribucion]]). La [[fig:muestras]] muestra ejemplos de cada clase."))
    add(("tab", "clases", "Clases de Fashion-MNIST: nombre usado en este trabajo y nombre original "
         "(Xiao, Rasul y Vollgraf, 2017, Tabla 2).",
         ["N.º", "Nombre en español", "Nombre original (inglés)"],
         [[c, D["clases"][c], D["clases_originales"][c]] for c in sorted(D["clases"], key=int)],
         [1.6, 5.0, 5.4], None, 10, None))
    add(("fig", "muestras", FIGURAS / "01_muestras.png", 8.5, 12.0,
         "Seis ejemplos por clase del conjunto de entrenamiento, elegidos al azar con semilla fija. "
         "Visualización invertida (prenda oscura sobre fondo claro); el modelo recibe la imagen original, con la "
         "prenda clara sobre fondo negro."))
    add(("fig", "distribucion", FIGURAS / "02_distribucion.png", 13.0, 6.5,
         "Cantidad de imágenes por clase en los conjuntos de entrenamiento y prueba."))
    add(("caja", "ATENCIÓN - los archivos CSV de la descarga no son la partición oficial", [
        "La descarga de Kaggle trae el conjunto en formato CSV y en formato idx. Al cotejarlos imagen por imagen "
        f"se comprobó que los CSV no reproducen la partición oficial: su conjunto de entrenamiento incluye "
        f"{n(diag['csv_train_en_test_oficial'])} imágenes del conjunto de prueba oficial y "
        f"{n(diag['csv_train_ausentes_del_oficial'])} ausentes de él, y comparte "
        f"{n(diag['imagenes_compartidas_csv_train_y_csv_test'])} imágenes con su propio conjunto de prueba.",
        "Mediante sumas MD5 se verificó que los archivos idx son idénticos a los del repositorio oficial, por lo "
        "que se trabajó con ellos: solo con la partición oficial los resultados son comparables con el "
        "benchmark. El detalle está en el [[anx:verificacion]]."]))
    add(("h2", "Preprocesamiento"))
    add(("p", "El preprocesamiento tuvo tres pasos: dividir los píxeles por 255 para llevarlos al intervalo "
         f"[0, 1]; dar a cada imagen la forma {dim(*D['forma_imagen'])}, con un único canal de intensidad, y "
         "codificar las etiquetas en formato one-hot con to_categorical, como en el ejemplo de la cátedra, de "
         f"modo que cada clase queda representada por un vector de {D['num_clases']} posiciones."))

    # 3. Consigna b
    add(("h1", "3. Consigna b: construcción de la red", {}))
    add(("p", "La red es un modelo secuencial (Sequential) de tensorflow.keras que sigue la red convolucional "
         "sobre MNIST del video de la cátedra, con un bloque convolucional más ([[fig:arquitectura]] y "
         "[[tab:arquitectura]])."))
    add(("h2", "Rol de cada capa"))
    add(("p", f"La entrada es una imagen de {dim(lado, lado)} píxeles con un canal. En cada capa de convolución, "
         f"un filtro o kernel de {kernel} recorre la imagen: en cada posición multiplica elemento a elemento sus "
         "pesos por la región que cubre, suma los productos y agrega un sesgo. Ese valor es una neurona del mapa "
         "de características; como el mismo filtro recorre toda la imagen, sus pesos se comparten. Las "
         "convoluciones usan la función de activación ReLU, f(x) = max{0, x}, que anula los valores negativos "
         "(material de cátedra: PPT Redes Neuronales Convolucionales)."))
    add(("p", f"Cada bloque termina con una agrupación máxima (max pooling) de {dim(*pool)}, que conserva el "
         f"máximo de cada grupo de {palabra(pool[0] * pool[1])} neuronas y reduce a la mitad el alto y el ancho; "
         f"en pool2, la entrada de {dim(*capas['conv2_2']['forma_salida'][:2])} queda en "
         f"{dim(*capas['pool2']['forma_salida'][:2])} porque se descartan la última fila y la última columna, "
         f"impares. La capa flatten convierte esos mapas ({forma('pool2')}) en un vector 1D de "
         f"{n(capas['flatten']['forma_salida'][0])} valores, que alimenta una capa totalmente conectada de "
         f"{cfg['densa']['unidades']} neuronas, donde se realiza el razonamiento de alto nivel. El dropout "
         "descarta al azar una fracción de las activaciones para reducir el sobreajuste; actúa solo durante el "
         "entrenamiento, y en la predicción se usan todas las neuronas. La salida tiene "
         f"{cfg['salida']['unidades']} neuronas con softmax, que entrega una probabilidad por clase."))
    add(("h2", "Decisiones de diseño"))
    add(("p", "El segundo bloque opera sobre mapas ya reducidos y combina características de mayor nivel. Los "
         f"filtros crecen de {cfg['conv1_1']['filtros']} a {cfg['conv2_1']['filtros']} mientras el pooling reduce "
         "la resolución, lo que suma capacidad sin multiplicar el costo. La primera convolución de cada bloque "
         "usa padding 'same' para conservar el tamaño y la información del borde; la segunda recorta un píxel "
         f"por lado. El dropout ({pct(cfg['drop1']['tasa'], 0)} tras cada bloque y {pct(cfg['drop3']['tasa'], 0)} "
         "tras la capa densa) replica el ejemplo de la cátedra. En Keras 3 la forma de entrada se declara con "
         "keras.Input, que reemplaza al argumento input_shape que usa el video."))
    add(("h2", "Función de pérdida y optimizador"))
    add(("p", "El ejemplo XOR del video usa error cuadrático medio con una única neurona sigmoide, adecuada "
         f"para una salida binaria. Aquí hay {D['num_clases']} clases excluyentes: softmax produce probabilidades "
         "que suman 1, y la entropía cruzada categórica penaliza la baja probabilidad asignada a la clase "
         "correcta, según la etiqueta one-hot. El optimizador es Adam."))
    add(("h2", "Parámetros"))
    vp = A["verificacion_parametros"]
    assert vp["conv1_1"]["coincide"] and vp["densa"]["coincide"]
    c11, dn = cfg["conv1_1"], capas["densa"]
    canales = capas["entrada"]["forma_salida"][-1]
    entradas = capas["flatten"]["forma_salida"][0]
    pf = A["prueba_forward"]
    add(("p", f"La red tiene {n(A['parametros_entrenables'])} parámetros, todos entrenables. La capa densa "
         f"posterior al flatten concentra {n(dn['parametros'])} ({pct(A['capa_con_mas_parametros']['porcentaje_del_total'] / 100)}); "
         f"las {palabra(len(convs))} convoluciones suman {n(A['parametros_por_tipo']['convolución']['parametros'])} "
         f"({pct(A['parametros_por_tipo']['convolución']['porcentaje_del_total'] / 100)}), gracias a los pesos "
         "compartidos. El cálculo se verificó a mano con (kh·kw·canales + 1)·filtros = "
         f"({c11['kernel'][0]}·{c11['kernel'][1]}·{canales} + 1)·{c11['filtros']} = {n(vp['conv1_1']['calculado'])} "
         f"para conv1_1, y con (entradas + 1)·neuronas = ({n(entradas)} + 1)·{cfg['densa']['unidades']} = "
         f"{n(vp['densa']['calculado'])} para la capa densa. Una prueba sin entrenar sobre "
         f"{palabra(pf['forma_salida'][0], False)} imágenes devolvió una salida de {dim(*pf['forma_salida'])} "
         "cuyas filas suman 1."))
    add(("tab", "arquitectura", "Arquitectura de la red: forma de salida y parámetros de cada capa.",
         ["Capa", "Tipo", "Forma de salida", "Parámetros"],
         [[c["nombre"], c["tipo"], dim(*c["forma_salida"]), n(c["parametros"])] for c in A["capas"]]
         + [["Total", "", "", n(A["total_parametros"])]],
         [3.0, 3.6, 4.2, 3.0], ["l", "l", "c", "r"], 10, {len(A["capas"]): {"negrita": True}}))
    add(("fig", "arquitectura", FIGURAS / "03_arquitectura.png", 16.0, 7.0,
         "Arquitectura de la red: una caja por capa con su nombre, tipo y forma de salida; el color indica el "
         "tipo de capa."))
    add(("p", "La función que construye y compila la red (src/modelo.py) es la siguiente:", {"juntar": True}))
    add(("code", fuente_python(SRC / "modelo.py", "construir_modelo")))

    # 4. Consigna c
    add(("h1", "4. Consigna c: proceso de aprendizaje", {}))
    add(("h2", "Aprendizaje supervisado y retropropagación"))
    add(("p", "El aprendizaje es supervisado: se muestran a la red patrones de entrada junto con su salida "
         "deseada, y una fórmula de minimización del error ajusta los pesos para acercar ambas (material de "
         "cátedra: apunte Aprendizaje en las Redes Neuronales). Los pesos de filtros y capas densas parten de "
         "valores aleatorios. Para cada lote, la red calcula su salida, mide el error con la entropía cruzada y, "
         "mediante retropropagación (backpropagation), obtiene las derivadas parciales de ese error respecto de "
         f"cada peso. Con ese gradiente, el optimizador Adam (variante del descenso de gradiente, con tasa de "
         f"aprendizaje {n(hp['learning_rate'], 3)}) ajusta pesos y filtros para reducir el error."))
    add(("h2", "Épocas, lotes e iteraciones"))
    add(("p", f"Los {n(part['n_entrenamiento'])} ejemplos de entrenamiento se procesan en lotes (batch) de "
         f"{hp['batch_size']}, mezclados en cada época; cada lote produce una iteración, es decir, un ajuste de "
         f"pesos. Una época, que es una pasada completa por los datos, tiene así {n(T['pasos_por_epoca'])} "
         f"iteraciones: {n(T['pasos_por_epoca'] - 1)} lotes completos y uno de {n(T['tamano_ultimo_lote'])}."))
    add(("h2", "Validación y early stopping"))
    por_clase_val = set(part["validacion_por_clase"].values())
    assert len(por_clase_val) == 1
    peores = previa["epocas_con_val_loss_peor_que_su_minimo_previo"]
    racha, maxima = 1, 1  # racha más larga de épocas consecutivas sin mejora en la corrida previa
    for anterior, siguiente in zip(peores, peores[1:]):
        racha = racha + 1 if siguiente == anterior + 1 else 1
        maxima = max(maxima, racha)
    assert maxima < previa["paciencia"] and not previa["corte_por_early_stopping"]
    add(("p", f"Del entrenamiento oficial se separaron, de forma estratificada, {n(part['n_validacion'])} "
         f"imágenes ({n(por_clase_val.pop())} por clase) para validación: la red no aprende de ellas; se usan "
         "para medir, al final de cada época, si lo aprendido generaliza. El early stopping vigila la pérdida de validación y, si no mejora en "
         f"{palabra(hp['early_stopping']['paciencia'], False)} épocas seguidas, detiene el entrenamiento y "
         "restaura los pesos de la mejor época. En una primera corrida, con un tope de "
         f"{previa['epocas_maximas']} épocas y paciencia {previa['paciencia']}, la pérdida de validación oscilaba "
         "sin tendencia clara: empeoró respecto de su mínimo en las épocas " + ", ".join(map(str, peores[:-1]))
         + f" y {peores[-1]}, pero nunca {palabra(previa['paciencia'], False)} veces seguidas, de modo que el early "
         f"stopping no llegó a actuar y la mejor época coincidió con el tope. Por eso se amplió el máximo a "
         f"{hp['epocas_maximas']} y la paciencia a {hp['early_stopping']['paciencia']}: así el corte responde a "
         "un estancamiento sostenido y no a una oscilación pasajera. El conjunto de prueba oficial no intervino "
         "en esta etapa ([[tab:hiper]])."))
    add(("tab", "hiper", "Hiperparámetros del proceso de aprendizaje.", ["Parámetro", "Valor"], [
        ["Optimizador", f"Adam (tasa de aprendizaje {n(hp['learning_rate'], 3)}, valor por defecto)"],
        ["Función de pérdida", "Entropía cruzada categórica"],
        ["Tamaño de lote", n(hp["batch_size"])],
        ["Épocas máximas", n(hp["epocas_maximas"])],
        ["Early stopping", f"pérdida de validación, paciencia {hp['early_stopping']['paciencia']}, "
                           "restaura los mejores pesos"],
        ["Mezcla de los datos", "en cada época"],
        ["Partición", f"{n(part['n_entrenamiento'])} de entrenamiento y {n(part['n_validacion'])} de validación "
                      "(estratificada)"],
        ["Iteraciones por época", f"{n(T['pasos_por_epoca'])} (último lote de {n(T['tamano_ultimo_lote'])})"],
        ["Semilla", n(R["dataset"]["semilla"])],
        ["Épocas ejecutadas / mejor época", f"{T['epocas_ejecutadas']} / {mejor}"],
    ], [5.2, 9.8], ["l", "l"], 10, None))
    add(("p", "Los callbacks y la llamada a fit (src/entrenar.py) son los siguientes:", {"juntar": True}))
    add(("code", fragmento_fit(SRC / "entrenar.py")))
    add(("h2", "Evolución del entrenamiento"))
    i = mejor - 1
    add(("p", f"La exactitud (accuracy) de validación pasó de {pct(H['val_accuracy'][0])} en la época 1 a "
         f"{pct(H['val_accuracy'][i])} en la {mejor}, y su pérdida de {n(H['val_loss'][0], 3)} a "
         f"{n(H['val_loss'][i], 3)} ([[fig:curva_precision]] y [[fig:curva_perdida]]). Hasta esa época, la "
         "exactitud de entrenamiento queda por debajo, porque se promedia durante la época con dropout activo. "
         f"Después las curvas se separan: la pérdida de entrenamiento sigue bajando (de {n(H['loss'][i], 3)} a "
         f"{n(H['loss'][-1], 3)}) y la de validación no vuelve a su mínimo. Es el comienzo del sobreajuste: el "
         f"early stopping cortó en la época {T['epocas_ejecutadas']} y restauró los pesos de la {mejor}. El "
         f"entrenamiento demoró {n(T['tiempo']['total_fit_min'], 2)} minutos en CPU, "
         f"{n(T['tiempo']['promedio_s'], 2)} s por época en promedio (detalle por época en la [[tab:epocas]], [[anx:epocas]])."))
    add(("fig", "curva_precision", FIGURAS / "05_curva_precision.png", 12.5, 6.6,
         "Exactitud (accuracy) de entrenamiento y validación por época. La línea punteada marca la mejor época (mínima "
         "pérdida de validación, cuyos pesos se restauran) y la discontinua, la época en que se detuvo el "
         "entrenamiento."))
    add(("fig", "curva_perdida", FIGURAS / "06_curva_perdida.png", 12.5, 6.6,
         "Pérdida (entropía cruzada) de entrenamiento y validación por época, con las mismas marcas que la "
         "[[fig:curva_precision]]."))
    add(("h2", "Mapas de activación"))
    f4 = T["figura_04"]
    add(("p", f"La [[fig:activaciones]] muestra, para la imagen de prueba n.º {f4['indice_test']} "
         f"({f4['clase'].lower()}), algunos mapas de la primera y de la última convolución del modelo entrenado. "
         f"Los de conv1_1 ({dim(*capas['conv1_1']['forma_salida'][:2])}) conservan la silueta y resaltan bordes "
         f"y contornos; los de conv2_2 ({dim(*capas['conv2_2']['forma_salida'][:2])}) concentran la activación en "
         "pocas zonas, como la franja horizontal de la suela: patrones más abstractos y de menor resolución, "
         "que la capa totalmente conectada combina para decidir la clase."))
    add(("fig", "activaciones", FIGURAS / "04_activaciones.png", 16.0, 6.0,
         f"Imagen de prueba n.º {f4['indice_test']} ({f4['clase'].lower()}) y "
         f"{palabra(len(f4['filtros_mostrados']['conv1_1']), False)} mapas de activación de conv1_1 "
         f"({cfg['conv1_1']['filtros']} filtros) y de conv2_2 ({cfg['conv2_2']['filtros']} filtros); se muestran "
         "los de mayor activación media. Visualización invertida de la entrada (prenda oscura sobre fondo "
         "claro); el modelo recibe la imagen original, con la prenda clara sobre fondo negro. Cada mapa usa su "
         "propia escala de color (más claro = mayor activación)."))

    # 5. Consigna d
    add(("h1", "5. Consigna d: validación del modelo", {}))
    add(("p", f"Recién en esta etapa se usó el conjunto de prueba oficial: {n(g['n_test'])} imágenes que la red "
         f"nunca vio. El modelo guardado obtuvo una exactitud de {pct(g['accuracy_test'])} y una pérdida de "
         f"{n(g['loss_test'], 3)}, frente a {pct(g['accuracy_validacion_mejor_epoca'])} y "
         f"{n(g['loss_validacion_mejor_epoca'], 3)} en validación: {n(abs(g['diferencia_accuracy_pp']), 2)} puntos "
         "porcentuales menos. La diferencia es chica, por lo que el modelo generaliza a datos nuevos; una leve "
         "caída es esperable, porque la mejor época se eligió mirando la validación."))
    add(("h2", "Métricas por clase"))
    pc = V["por_clase"]
    mj, pj = V["mejor_clase_f1"], V["peor_clase_f1"]
    add(("p", f"La [[tab:metricas]] resume la precisión (precision: qué parte de lo predicho como una clase es correcto), "
         "el recall (qué parte de cada clase se reconoce) y el f1-score, su media armónica. La mejor clase es "
         f"{mj['clase']} (f1 = {n(mj['f1'], 3)}) y la peor, {pj['clase']} (f1 = {n(pj['f1'], 3)}; precisión "
         f"{n(pc[pj['clase']]['precision'], 3)} y recall {n(pc[pj['clase']]['recall'], 3)}). Los promedios macro y "
         f"ponderado coinciden en f1 = {n(V['macro_avg']['f1-score'], 3)}, porque las clases tienen el mismo "
         "soporte."))
    assert round(V["macro_avg"]["f1-score"], 3) == round(V["weighted_avg"]["f1-score"], 3)
    filas_m = [[c, n(v["precision"], 3), n(v["recall"], 3), n(v["f1-score"], 3), n(v["support"])]
               for c, v in pc.items()]
    filas_m += [["Promedio macro", n(V["macro_avg"]["precision"], 3), n(V["macro_avg"]["recall"], 3),
                 n(V["macro_avg"]["f1-score"], 3), n(V["macro_avg"]["support"])],
                ["Promedio ponderado", n(V["weighted_avg"]["precision"], 3), n(V["weighted_avg"]["recall"], 3),
                 n(V["weighted_avg"]["f1-score"], 3), n(V["weighted_avg"]["support"])],
                ["Exactitud (accuracy)", "", "", n(V["accuracy"], 3), n(g["n_test"])]]
    nombres_pc = list(pc)
    add(("tab", "metricas", "Métricas por clase sobre el conjunto de prueba (classification_report de "
         "scikit-learn). Se resaltan la mejor y la peor clase según f1.",
         ["Clase", "Precisión", "Recall", "F1-score", "Soporte"], filas_m, [4.6, 2.6, 2.6, 2.6, 2.4],
         ["l", "c", "c", "c", "c"], 10,
         {nombres_pc.index(mj["clase"]): {"negrita": True, "fondo": "E2EFDA"},
          nombres_pc.index(pj["clase"]): {"negrita": True, "fondo": "FBE5D6"},
          len(pc): {"cursiva": True}, len(pc) + 1: {"cursiva": True}, len(pc) + 2: {"negrita": True}}))
    add(("h2", "Matriz de confusión"))
    pares = V["pares_mas_confundidos"]
    tor = V["confusion_prendas_de_torso"]
    add(("p", "En la [[fig:confusion]], cada fila es la clase real y cada columna, la predicha. Los pares más "
         "confundidos son " + ", ".join(f"{p['real']} → {p['predicha']} ({n(p['cantidad'])} casos)" for p in pares[:-1])
         + f" y {pares[-1]['real']} → {pares[-1]['predicha']} ({n(pares[-1]['cantidad'])}). Los "
         f"{palabra(len(pares), False)} caen dentro del grupo de prendas de torso ("
         + ", ".join(x.lower() for x in tor["grupo"][:-1]) + f" y {tor['grupo'][-1].lower()}), que reúne "
         f"{n(tor['errores_dentro_del_grupo'])} de los {n(tor['errores_totales'])} errores "
         f"({pct(tor['porcentaje_de_los_errores'] / 100)}): son siluetas parecidas que, a {dim(lado, lado)} "
         "píxeles, se distinguen por detalles como botones, cuello o largo de manga."))
    add(("fig", "confusion", FIGURAS / "07_matriz_confusion.png", 11.5, 10.5,
         f"Matriz de confusión sobre las {n(g['n_test'])} imágenes de prueba: filas, clase real; columnas, "
         "clase predicha."))
    add(("h2", "Ejemplos cotejados"))
    por_clase_ej = ej["total"] // D["num_clases"]
    assert por_clase_ej * D["num_clases"] == ej["total"]
    fallos = [e for e in ej["detalle"] if not e["acierta"]]
    segundas = {e["top2"][1]["clase"] for e in fallos}
    grupos = {}
    for e in fallos:
        clave = (e["clase_real"].lower(), e["clase_predicha"].lower())
        grupos.setdefault(clave, []).append(e["probabilidad_predicha"])
    texto_fallos = "; ".join(
        (f"{palabra(len(ps), False)} casos de {real} predicho como {pred}" if len(ps) > 1
         else f"{real} predicho como {pred}") + " (p = " + " y ".join(n(x, 2) for x in ps) + ")"
        for (real, pred), ps in grupos.items())
    add(("p", f"Se eligieron {palabra(ej['total'], False)} imágenes de prueba al azar antes de ver las "
         f"predicciones ({palabra(por_clase_ej, False)} por clase, con semilla fija) y se cotejó la clase predicha con la real "
         f"([[fig:ejemplos]]). El modelo acertó {ej['aciertos']}, en línea con los "
         f"{n(ej['aciertos_esperados_segun_accuracy_global'], 2)} aciertos esperables según la exactitud global. "
         f"Los fallos fueron: {texto_fallos}"
         + (f", siempre con {segundas.pop()} como segunda opción" if len(fallos) > 1 and len(segundas) == 1 else "")
         + ". El modelo duda, con probabilidades bajas, justamente entre prendas de torso. El detalle de cada "
         "ejemplo, con su top-2, está en el [[anx:ejemplos]] ([[tab:ejemplos]])."))
    assert all(e["clase_real"] in tor["grupo"] and e["clase_predicha"] in tor["grupo"] for e in fallos)
    add(("fig", "ejemplos", FIGURAS / "08_ejemplos_cotejados.png", 14.0, 18.0,
         f"{palabra(ej['total'], False).capitalize()} ejemplos de prueba ({palabra(por_clase_ej, False)} por clase, al azar con semilla fija) "
         "con su clase real, la predicha y la probabilidad asignada. Verde y ✓: acierto; rojo y ✗: error. "
         "Visualización invertida (prenda oscura sobre fondo claro); el modelo recibe la imagen original, con la "
         "prenda clara sobre fondo negro."))
    add(("h2", "Errores con alta confianza"))
    clasif = E["clasificacion"]
    assert sorted(map(int, clasif)) == sorted(err["indices"]), "el análisis cualitativo no corresponde a los errores"
    cuenta = {k: sum(1 for v in clasif.values() if v["categoria"] == k) for k in E["categorias"]}
    det = {str(e["indice"]): e for e in err["detalle"]}
    largas = [det[i]["clase_real"].lower() for i, v in clasif.items() if v["categoria"] == "prenda_larga"]
    fallas = [v["nota"] for v in clasif.values() if v["categoria"] == "falla"]
    pmin = min(e["probabilidad_predicha"] for e in err["detalle"])
    add(("fig", "errores", FIGURAS / "09_errores.png", 14.0, 18.0,
         f"Los {palabra(len(err['indices']), False)} errores del conjunto de prueba con mayor probabilidad asignada "
         "a la clase equivocada, con su top-3. Visualización invertida (prenda oscura sobre fondo claro); el "
         "modelo recibe la imagen original, con la prenda clara sobre fondo negro."))
    add(("p", f"De los {n(err['errores_totales'])} errores del conjunto de prueba, la [[fig:errores]] muestra los "
         f"{palabra(len(err['indices']), False)} con mayor confianza: todos asignan al menos {n(math.floor(pmin * 1000) / 1000, 3)} de "
         f"probabilidad a la clase equivocada. A simple vista, la mayoría son entendibles: "
         f"{palabra(cuenta['calzado'], False)} son calzados con recortes o plataformas, ambiguos entre bota, "
         f"sandalia y zapatilla; {palabra(cuenta['prenda_larga'], False)} son prendas largas y angostas ("
         + ", ".join(cantidad_prenda(largas.count(x), x) for x in sorted(set(largas))) + ") leídas como vestido, y "
         f"{palabra(cuenta['chaqueta_botones'], False)} es una chaqueta estampada con botones tomada por camisa. "
         f"Los otros {palabra(cuenta['falla'], False)} parecen fallas del modelo o imágenes atípicas: "
         + "; ".join(fallas) + ". Que el modelo se equivoque con probabilidades cercanas a 1 muestra que la "
         "salida softmax no equivale a certeza."))
    add(("h2", "Aporte de la convolución y comparación con el benchmark"))
    cnn_row, mlp_row = comp["tabla"]
    refs = {r["modelo"]: r for r in B["referencias"]}
    svc = next(r for r in B["referencias"] if r["modelo"].startswith("SVC"))
    cnn_debajo = [r for r in B["referencias"] if "Conv" in r["modelo"] and r["accuracy_test"] < g["accuracy_test"]]
    cnn_encima = [r for r in B["referencias"] if "Conv" in r["modelo"] and r["accuracy_test"] > g["accuracy_test"]]
    mlp_ref = next(r for r in B["referencias"] if r["modelo"].startswith("MLP"))
    assert svc["accuracy_test"] < g["accuracy_test"] and len(cnn_encima) == 1
    add(("p", "Para medir qué aporta la convolución se entrenó, en idénticas condiciones, una red sin "
         "convolución: un perceptrón multicapa que aplana la imagen y la pasa por una capa densa de "
         f"{cfg['densa']['unidades']} neuronas. Alcanzó {pct(mlp_row['accuracy_test'])} con "
         f"{n(mlp_row['parametros'])} parámetros, frente a {pct(cnn_row['accuracy_test'])} de la CNN con "
         f"{n(cnn_row['parametros'])}: la convolución aporta {n(comp['ventaja_cnn_pp'], 2)} puntos porcentuales "
         "([[tab:comparacion]]). El resultado supera al mejor clasificador clásico del artículo original, un SVC con "
         f"kernel RBF ({pct_fuente(svc['accuracy_test'])}; Xiao, Rasul y Vollgraf, 2017, Tabla 3), y a "
         f"{palabra(len(cnn_debajo), False)} CNN chicas de la tabla de contribuciones del repositorio oficial, "
         + " y ".join(f"«{r['modelo']}» ({pct_fuente(r['accuracy_test'])})" for r in cnn_debajo)
         + f". Queda por debajo de «{cnn_encima[0]['modelo']}» ({pct_fuente(cnn_encima[0]['accuracy_test'])}), "
         f"que agrega normalización por lotes. El MLP propio queda cerca del «{mlp_ref['modelo']}» informado allí "
         f"({pct_fuente(mlp_ref['accuracy_test'])}). Las cifras de esa tabla las informan sus autores, en "
         "condiciones no homogéneas (repositorio oficial de Fashion-MNIST)."))
    filas_c = []
    resaltar_c = {}
    for k, f in enumerate(bench["tabla_ordenada"]):
        if f["fuente"] == "este trabajo":
            propia = cnn_row if f["modelo"].startswith("CNN") else mlp_row
            filas_c.append([f["modelo"], n(propia["parametros"]), pct(f["accuracy_test"]),
                            f"{n(propia['tiempo_entrenamiento_s'], 1)} s ({propia['epocas']} épocas)", "Este trabajo"])
            resaltar_c[k] = {"negrita": True, "fondo": "E2EFDA"} if f["modelo"].startswith("CNN") else {"negrita": True}
        else:
            r = refs[f["modelo"]]
            fuente = ("Xiao, Rasul y Vollgraf (2017), Tabla 3" if r is svc
                      else "Repositorio oficial, tabla de contribuciones")
            filas_c.append([f["modelo"].split(" — ")[0], "s/d", pct_fuente(f["accuracy_test"]), "s/d", fuente])
    add(("tab", "comparacion", "Comparación de la CNN con la red sin convolución y con resultados publicados "
         "(exactitud sobre el conjunto de prueba oficial; s/d: sin dato en la fuente).",
         ["Modelo", "Parámetros", "Exactitud", "Entrenamiento", "Fuente"], filas_c,
         [4.4, 2.1, 2.0, 3.0, 4.5], ["l", "c", "c", "c", "l"], 9, resaltar_c))
    add(("h2", "Predicción sobre imágenes nuevas"))
    pr = V["prueba_predecir"]
    assert all(x["coincide"] for x in pr)
    add(("p", "El script predecir.py permite usar el modelo con imágenes nuevas: las pasa a escala de grises, "
         f"las lleva a {dim(lado, lado)} píxeles con antialiasing, las invierte si el fondo es claro y devuelve la "
         f"clase con su top-3. Sobre {palabra(len(pr), False)} imágenes de prueba exportadas como PNG, una "
         "original y otra invertida, reprodujo exactamente las predicciones de la validación, lo que confirma "
         "que el preprocesamiento es el mismo."))
    add(("caja", "RESULTADO", [
        f"La CNN clasifica correctamente {n(g['aciertos'])} de las {n(g['n_test'])} imágenes de prueba "
        f"({pct(g['accuracy_test'])}), con una diferencia de {n(g['diferencia_accuracy_pp'], 2)} puntos "
        f"porcentuales respecto de la validación y {n(comp['ventaja_cnn_pp'], 2)} puntos por encima de la red sin "
        "convolución.",
        f"Sus errores se concentran en las prendas de torso ({pct(tor['porcentaje_de_los_errores'] / 100)} del "
        f"total) y la clase más difícil es {pj['clase']} (f1 = {n(pj['f1'], 3)})."]))

    # 6. Conclusiones
    add(("h1", "6. Conclusiones", {}))
    posicion = len([r for r in B["referencias"] if r["accuracy_test"] > g["accuracy_test"]])
    add(("concl", f"El modelo alcanzó una exactitud de {pct(g['accuracy_test'])} sobre las {n(g['n_test'])} "
         f"imágenes de prueba, {n(abs(g['diferencia_accuracy_pp']), 2)} puntos porcentuales por debajo de la "
         f"validación, por lo que generaliza a datos que nunca vio. Supera en {n(comp['ventaja_cnn_pp'], 2)} "
         "puntos a la red sin convolución entrenada en idénticas condiciones: los filtros aprovechan la "
         "estructura espacial de la imagen, que el perceptrón multicapa pierde al aplanarla. Frente al benchmark "
         f"queda por encima de {palabra(len(B['referencias']) - posicion, False)} de las "
         f"{palabra(len(B['referencias']), False)} referencias citadas, incluido el mejor clasificador clásico "
         f"del artículo original ({pct_fuente(svc['accuracy_test'])}), y por debajo de la variante con normalización por "
         f"lotes ({pct_fuente(cnn_encima[0]['accuracy_test'])})."))
    add(("concl", f"**Limitaciones.** Los errores se concentran en las prendas de torso "
         f"({pct(tor['porcentaje_de_los_errores'] / 100)} del total), y {pj['clase']} es la clase más débil "
         f"(f1 = {n(pj['f1'], 3)}). El modelo aprendió con imágenes de {dim(lado, lado)} píxeles, centradas y con "
         "la prenda clara sobre fondo negro, de modo que no generaliza a fotografías reales sin un "
         "preprocesamiento equivalente (recorte, centrado, escala de grises e inversión); predecir.py cubre solo "
         "una parte de ese proceso. Además, el entrenamiento se hizo solo en CPU: cada corrida demora "
         f"{n(T['tiempo']['total_fit_min'], 2)} minutos, lo que encarece explorar variantes."))
    add(("concl", "**Mejoras posibles.** Incorporar normalización por lotes, que en el benchmark lleva una "
         f"arquitectura de dos convoluciones a {pct_fuente(cnn_encima[0]['accuracy_test'])}; aplicar aumento de "
         "datos (data augmentation), con desplazamientos y volteos horizontales, para exponer más variantes de "
         "las siluetas de torso que hoy se confunden, y sumar filtros al segundo bloque: hoy la capa densa "
         f"concentra el {pct(A['capa_con_mas_parametros']['porcentaje_del_total'] / 100)} de los parámetros, "
         "mientras que las convoluciones, que extraen los detalles que distinguen botones, cuellos o mangas, "
         "tienen relativamente poca capacidad."))

    # Anexo A
    add(("h1", "[[anx:entorno]]. Entorno y ejecución", {"anexo": "entorno"}))
    hw = T["hardware"]
    nombres = {"python": "Python", "tensorflow": "TensorFlow", "keras": "Keras", "numpy": "NumPy",
               "pandas": "pandas", "scikit-learn": "scikit-learn", "matplotlib": "Matplotlib"}
    filas_e = [[nombres.get(k.strip(), k.strip()), v.strip()]
               for k, v in (l.split(":", 1) for l in entorno if ":" in l)]
    filas_e += [["Procesador", hw["procesador"]],
                ["Núcleos", f"{hw['nucleos_fisicos']} físicos, {hw['nucleos_logicos']} lógicos"],
                ["Memoria RAM", f"{n(hw['ram_gb'], 1)} GB"], ["GPU", hw["gpu"]]]
    add(("tab", "entorno", "Versiones del entorno y hardware usados en todas las corridas.",
         ["Componente", "Versión o detalle"], filas_e, [4.5, 10.5], ["l", "l"], 10, None))
    version_tf = next(v for k, v in filas_e if k.lower() == "tensorflow")
    version_keras = next(v for k, v in filas_e if k.lower() == "keras")
    add(("p", "La [[tab:entorno]] resume las versiones y el hardware de todas las corridas. El video de la cátedra usa Python 3.6 con el paquete keras independiente. Aquí se usa "
         f"tensorflow.keras (TensorFlow {version_tf}, Keras {version_keras}), que mantiene la misma API: "
         "Sequential, Conv2D, MaxPooling2D, Dropout, Flatten, Dense, to_categorical, fit, evaluate, predict y el "
         "grabado del modelo. El único cambio visible es la declaración de la entrada con keras.Input. La "
         f"semilla es {n(R['dataset']['semilla'])} en random, NumPy y TensorFlow; sin determinismo forzado de "
         "operaciones, en CPU puede haber variaciones mínimas entre corridas."))
    add(("p", "Comandos para ejecutar cada etapa desde la carpeta tp_cnn:", {"juntar": True}))
    add(("code", "\n".join([
        "python -m venv .venv",
        ".venv\\Scripts\\activate            (Linux/Mac: source .venv/bin/activate)",
        "pip install -r requirements.txt",
        "python src/datos.py                 # a. carga, verificación y figuras del dataset",
        "python src/modelo.py                # b. construcción de la red (sin entrenar)",
        "python src/entrenar.py --prueba     # mide una época sin guardar nada",
        "python src/entrenar.py              # c. entrenamiento completo",
        "python src/mlp_referencia.py        # red sin convolución para comparar",
        "python src/validar.py               # d. validación sobre el test oficial",
        "python src/predecir.py imagen.png   # predicción sobre una imagen suelta",
        "python src/generar_informe.py --modulo N --fecha DD/MM/AAAA",
    ])))
    add(("p", "Estructura del proyecto (codigo_TP_CNN.zip incluye src/, requirements.txt, README.md, el modelo "
         "entrenado y el logo; outputs/ se genera al ejecutar los scripts, notebook/ y entrega/ son los entregables "
         "y archive/ es la descarga de Kaggle):", {"juntar": True}))
    add(("code", "\n".join([
        "tp_cnn/src/       datos, modelo, entrenar, mlp_referencia, validar, predecir, generar_informe",
        "tp_cnn/outputs/   resultados.json, historial.json, entorno.txt, modelo .keras, figuras/, pruebas/",
        "tp_cnn/notebook/  TP_CNN_FashionMNIST.ipynb (flujo completo, ejecutado)",
        "tp_cnn/entrega/   informe .docx y .pdf, notebook y código comprimido",
        "archive/          datos descargados de Kaggle (no se incluyen en la entrega)",
    ])))

    # Anexo: ejemplos cotejados
    add(("h1", "[[anx:ejemplos]]. Ejemplos cotejados", {"anexo": "ejemplos"}))
    add(("tab", "ejemplos", "Cotejo de los ejemplos: clase real, clase predicha, probabilidad y segunda opción.",
         ["N.º (prueba)", "Clase real", "Clase predicha", "p", "2.ª opción (p)", "Resultado"],
         [[n(e["indice"]), e["clase_real"], e["clase_predicha"], n(e["probabilidad_predicha"], 2),
           f"{e['top2'][1]['clase']} ({n(e['top2'][1]['probabilidad'], 2)})",
           ("✓", VERDE, SIMBOLOS) if e["acierta"] else ("✗", ROJO, SIMBOLOS)] for e in ej["detalle"]],
         [2.1, 2.9, 2.9, 1.3, 3.8, 2.4], ["c", "l", "l", "c", "l", "c"], 9,
         {k: {"fondo": "FBE5D6"} for k, e in enumerate(ej["detalle"]) if not e["acierta"]}))

    # Anexo: entrenamiento por época
    add(("h1", "[[anx:epocas]]. Entrenamiento por época", {"anexo": "epocas"}))

    filas_b, resaltar_b = [], {}
    for k in range(len(H["loss"])):
        etiqueta = str(k + 1)
        if k + 1 == mejor:
            etiqueta += " (mejor)"
            resaltar_b[k] = {"negrita": True, "fondo": "E2EFDA"}
        elif k + 1 == T["epocas_ejecutadas"]:
            etiqueta += " (detención)"
            resaltar_b[k] = {"cursiva": True, "fondo": GRIS_CAJA}
        filas_b.append([etiqueta, n(H["accuracy"][k], 4), n(H["val_accuracy"][k], 4), n(H["loss"][k], 4),
                        n(H["val_loss"][k], 4), n(H["segundos"][k], 1)])
    add(("tab", "epocas", "Métricas por época de la corrida definitiva (historial.json). Se resaltan la mejor "
         "época, cuyos pesos se restauran, y la de detención.",
         ["Época", "Exactitud entr.", "Exactitud val.", "Pérdida entr.", "Pérdida val.", "Segundos"], filas_b,
         [2.8, 2.5, 2.5, 2.5, 2.5, 2.0], ["l", "c", "c", "c", "c", "c"], 9, resaltar_b))

    # Anexo: verificación del conjunto de datos
    add(("h1", "[[anx:verificacion]]. Verificación del conjunto de datos", {"anexo": "verificacion"}))
    add(("p", "La descarga de Kaggle incluye el conjunto en formato CSV y en formato idx, y ambos se cotejaron "
         "imagen por imagen. Los CSV no reproducen la partición oficial: su entrenamiento contiene "
         f"{n(diag['csv_train_en_test_oficial'])} imágenes de la prueba oficial y "
         f"{n(diag['csv_train_ausentes_del_oficial'])} ausentes del conjunto oficial, su prueba contiene "
         f"{n(diag['csv_test_en_train_oficial'])} imágenes del entrenamiento oficial, y ambos comparten "
         f"{n(diag['imagenes_compartidas_csv_train_y_csv_test'])} imágenes. Los archivos idx usados tienen las "
         "sumas MD5 de la [[tab:md5]], que coinciden con las de los archivos del repositorio oficial "
         "descomprimidos, cuyos .gz tienen a su vez los MD5 que publica el repositorio. El diagnóstico completo "
         "está en resultados.json (sección dataset)."))
    add(("tab", "md5", "Sumas MD5 de los archivos idx usados.", ["Archivo", "MD5"],
         [[k, v] for k, v in D["md5_archivos_usados"].items()], [5.6, 8.4], ["l", "l"], 9, None))
    return b


# ---------------------------------------------------------------- render

def numerar(bloques):
    """Asigna 'Figura N', 'Tabla N' y la letra de cada anexo por orden de aparición."""
    refs, nf, nt, na = {}, 0, 0, 0
    for blq in bloques:
        if blq[0] == "h1" and blq[2].get("anexo"):
            refs[f"anx:{blq[2]['anexo']}"] = f"Anexo {chr(ord('A') + na)}"
            na += 1
        elif blq[0] == "fig":
            nf += 1
            refs[f"fig:{blq[1]}"] = f"Figura {nf}"
        elif blq[0] == "tab":
            nt += 1
            refs[f"tab:{blq[1]}"] = f"Tabla {nt}"
    return refs


def resolver(texto, refs):
    salida = re.sub(r"\[\[(fig|tab|anx):(\w+)\]\]", lambda m: refs[f"{m.group(1)}:{m.group(2)}"], texto)
    assert "[[" not in salida, salida
    return salida


def render(bloques, refs, doc):
    for blq in bloques:
        tipo = blq[0]
        if tipo == "h1":
            titulo1(doc, resolver(blq[1], refs), blq[2].get("salto", False))
        elif tipo == "h2":
            titulo2(doc, blq[1])
        elif tipo in ("p", "intro", "concl"):
            opciones = blq[2] if len(blq) > 2 else {}
            cuerpo(doc, resolver(blq[1], refs), juntar=opciones.get("juntar", False))
        elif tipo == "item":
            cuerpo(doc, blq[1], sangria=0.63)
        elif tipo == "fig":
            _, etiqueta, archivo, ancho, alto, epigrafe = blq
            figura(doc, archivo, refs[f"fig:{etiqueta}"], resolver(epigrafe, refs), ancho, alto)
        elif tipo == "tab":
            _, etiqueta, titulo, enc, filas, anchos, alin, tam, resaltar = blq
            tabla(doc, refs[f"tab:{etiqueta}"], resolver(titulo, refs), enc, filas, anchos, alin, tam, resaltar,
                  compacta=etiqueta in TABLAS_COMPACTAS)
        elif tipo == "code":
            codigo(doc, blq[1])
        elif tipo == "caja":
            caja(doc, blq[1], [resolver(t, refs) for t in blq[2]])


def contar_palabras(bloques, tipo, refs):
    texto = " ".join(resolver(b[1], refs) for b in bloques if b[0] == tipo)
    return len(re.sub(r"\*\*", "", texto).split())


def main():
    parser = argparse.ArgumentParser(description="Genera el informe .docx del TP")
    parser.add_argument("--modulo", default="[N]")
    parser.add_argument("--fecha", default="[DD/MM/AAAA]")
    parser.add_argument("--salida", default=str(SALIDA_POR_DEFECTO))
    args = parser.parse_args()

    R = json.loads((SALIDAS / "resultados.json").read_text(encoding="utf-8"))
    H = json.loads((SALIDAS / "historial.json").read_text(encoding="utf-8"))
    B = json.loads((SALIDAS / "benchmark_referencias.json").read_text(encoding="utf-8"))
    E = json.loads((SALIDAS / "analisis_errores_alta_confianza.json").read_text(encoding="utf-8"))
    entorno = (SALIDAS / "entorno.txt").read_text(encoding="utf-8").splitlines()
    entorno = [l for l in entorno if not l.startswith("sistema operativo")]

    bloques = construir_bloques(R, H, B, E, entorno, args)
    refs = numerar(bloques)
    palabras_intro, palabras_concl = contar_palabras(bloques, "intro", refs), contar_palabras(bloques, "concl", refs)
    assert palabras_intro <= 200, f"introducción: {palabras_intro} palabras"
    assert palabras_concl <= 300, f"conclusiones: {palabras_concl} palabras"

    doc = Document()
    configurar_documento(doc)
    portada(doc, args.modulo, "Redes Neuronales Convolucionales (CNN)", args.fecha)
    render(bloques, refs, doc)
    salida = Path(args.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(salida)

    mapa = {Path(b[2]).name: refs[f"fig:{b[1]}"] for b in bloques if b[0] == "fig"}
    print(json.dumps({"salida": str(salida), "figuras": mapa,
                      "tablas": {k: v for k, v in refs.items() if k.startswith("tab:")},
                      "palabras_introduccion": palabras_intro, "palabras_conclusiones": palabras_concl},
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
